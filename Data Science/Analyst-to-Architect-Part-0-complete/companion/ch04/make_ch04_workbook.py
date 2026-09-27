# Builds numbers_practice.xlsx for Chapter 4 from the one-year database (riverstone_2025).
# Run: python3 make_ch04_workbook.py   (needs PostgreSQL with riverstone_2025 loaded, and openpyxl)
import subprocess
from openpyxl import Workbook
from openpyxl.styles import Font
def q(sql):
    r = subprocess.run(['su','postgres','-c',"psql -X -q -At -F '|' -d riverstone_2025"], input=sql, capture_output=True, text=True, cwd='/tmp')
    return [l.split('|') for l in r.stdout.strip().splitlines()]
wb = Workbook(); b = Font(bold=True)
ws = wb.active; ws.title = 'monthly'
ws.append(['month', 'revenue', 'product_cost', 'target', 'gross_margin_pct', 'pct_change', 'margin_change_points', 'pct_of_target'])
rows = q("""SELECT to_char(m.mon,'YYYY-MM'), ROUND(m.rev,2), ROUND(m.cost,2), t.target_revenue
            FROM (SELECT DATE_TRUNC('month',order_date)::date AS mon, SUM(net_revenue) AS rev, SUM(product_cost) AS cost FROM sales_lines GROUP BY 1) m
            JOIN sales_targets t ON t.target_month=m.mon ORDER BY 1;""")
for i, (mo, rev, cost, tgt) in enumerate(rows, start=2):
    ws.append([mo, float(rev), float(cost), float(tgt), f'=(B{i}-C{i})/B{i}',
               None if i == 2 else f'=(B{i}-B{i-1})/B{i-1}',
               None if i == 2 else f'=(E{i}-E{i-1})*100', f'=B{i}/D{i}'])
ws['A15'] = 'total'; ws['B15'] = '=SUM(B2:B13)'; ws['D15'] = '=SUM(D2:D13)'; ws['H15'] = '=B15/D15'
ws['A17'] = 'average of monthly % changes'; ws['F17'] = '=AVERAGE(F3:F13)'
ws['A18'] = 'compound monthly growth, Jan to Dec'; ws['F18'] = '=(B13/B2)^(1/11)-1'
ws['A19'] = 'same, with RRI'; ws['F19'] = '=_xlfn.RRI(11,B2,B13)'
ws['A21'] = 'CAGR needed: 4,335,471 to 6,000,000 in 3 years'; ws['F21'] = '=(6000000/4335471)^(1/3)-1'
ws['A22'] = 'same, with RRI'; ws['F22'] = '=_xlfn.RRI(3,4335471,6000000)'
ws2 = wb.create_sheet('orders'); ws2.append(['order_id', 'order_value'])
for oid, v in q("SELECT order_id, ROUND(SUM(net_revenue),2) FROM sales_lines GROUP BY 1 ORDER BY 1;"):
    ws2.append([int(oid), float(v)])
n = ws2.max_row
ws2['D1'] = 'mean'; ws2['E1'] = f'=AVERAGE(B2:B{n})'
ws2['D2'] = 'median'; ws2['E2'] = f'=MEDIAN(B2:B{n})'
ws2['D3'] = 'orders above the mean'; ws2['E3'] = f'=COUNTIF(B2:B{n},">"&E1)'
ws2['D4'] = 'count'; ws2['E4'] = f'=COUNT(B2:B{n})'
ws3 = wb.create_sheet('discounts'); ws3.append(['discount_pct', 'order_lines', 'gross_value'])
for d, c, g in q("SELECT oi.discount_pct, COUNT(*), SUM(oi.quantity*oi.unit_price) FROM order_items oi JOIN orders o USING(order_id) WHERE o.status<>'Cancelled' GROUP BY 1 ORDER BY 1;"):
    ws3.append([float(d), int(c), float(g)])
ws3['E1'] = 'simple average discount (per line)'; ws3['F1'] = '=SUMPRODUCT(A2:A6,B2:B6)/SUM(B2:B6)'
ws3['E2'] = 'weighted by gross value'; ws3['F2'] = '=SUMPRODUCT(A2:A6,C2:C6)/SUM(C2:C6)'
ws3['E3'] = 'most common discount (mode of lines)'; ws3['F3'] = '=INDEX(A2:A6,MATCH(MAX(B2:B6),B2:B6,0))'
ws4 = wb.create_sheet('segments'); ws4.append(['segment', 'revenue', 'share', 'share_rounded'])
S = q("SELECT c.segment, ROUND(SUM(net_revenue),2) FROM sales_lines s JOIN customers c USING(customer_id) GROUP BY 1 ORDER BY 2 DESC;")
for i, (seg, r) in enumerate(S, start=2):
    ws4.append([seg, float(r), f'=B{i}/SUM($B$2:$B$4)', f'=ROUND(C{i}*100,0)'])
ws4['A5'] = 'total'; ws4['C5'] = '=SUM(C2:C4)'; ws4['D5'] = '=SUM(D2:D4)'
for w in wb.worksheets:
    for c in w[1]: c.font = b
wb.save('numbers_practice.xlsx'); print('saved', n-1, 'orders')
