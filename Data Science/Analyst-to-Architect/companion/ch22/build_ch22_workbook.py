"""
Analyst to Architect — Chapter 22: Statistics Without Fooling Yourself
build_ch22_workbook.py — builds ch22_by_hand.xlsx, the spreadsheet versions of the chapter's by-hand sums.

Run from this folder:  python3 build_ch22_workbook.py     (needs openpyxl)

Sheets:
  Interval     section 22.1: a survey proportion's 95% interval, and the margin for a mean
  SampleSize   section 22.3: the two z multipliers and the sample size for a 24% -> 27% open rate
  Line         section 22.10: six customers, SLOPE / INTERCEPT / RSQ / CORREL / FORECAST.LINEAR and residuals

The formulas are written as text; Excel, Google Sheets and LibreOffice calculate them when the file opens.
Functions newer than Excel 2007 carry the "_xlfn." prefix the xlsx format requires; Excel shows them without it.
"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = pathlib.Path(__file__).resolve().parent
BOLD = Font(bold=True)
wb = Workbook()

# ---- 1. Interval (section 22.1) ---------------------------------------------------------------
ws = wb.active
ws.title = "Interval"
rows = [
    ("A survey: 400 customers, 62% satisfied", None),
    ("n", 400),
    ("p-hat", 0.62),
    ("standard error", "=SQRT(B3*(1-B3)/B2)"),
    ("z for 95%", "=_xlfn.NORM.S.INV(0.975)"),
    ("margin", "=B5*B4"),
    ("lower", "=B3-B6"),
    ("upper", "=B3+B6"),
    (None, None),
    ("A mean: 200 orders, standard deviation 17,862", None),
    ("n", 200),
    ("standard deviation", 17862),
    ("margin, normal multiplier", "=_xlfn.CONFIDENCE.NORM(0.05,B12,B11)"),
    ("margin, t multiplier", "=_xlfn.CONFIDENCE.T(0.05,B12,B11)"),
    ("t multiplier", "=_xlfn.T.INV.2T(0.05,B11-1)"),
]
for r, (label, value) in enumerate(rows, start=1):
    ws.cell(r, 1, label)
    if value is not None:
        ws.cell(r, 2, value)
ws["A1"].font = ws["A10"].font = BOLD
ws.column_dimensions["A"].width = 44

# ---- 2. SampleSize (section 22.3) -------------------------------------------------------------
ws = wb.create_sheet("SampleSize")
rows = [
    ("Open rate 24% -> 27%, alpha 0.05 two-sided, power 0.80", None),
    ("p1", 0.24),
    ("p2", 0.27),
    ("p-bar", "=(B2+B3)/2"),
    ("z alpha/2", "=_xlfn.NORM.S.INV(0.975)"),
    ("z beta", "=_xlfn.NORM.S.INV(0.8)"),
    ("n per variant", "=ROUNDUP((B5*SQRT(2*B4*(1-B4))+B6*SQRT(B2*(1-B2)+B3*(1-B3)))^2/(B3-B2)^2,0)"),
]
for r, (label, value) in enumerate(rows, start=1):
    ws.cell(r, 1, label)
    if value is not None:
        ws.cell(r, 2, value)
ws["A1"].font = BOLD
ws.column_dimensions["A"].width = 52

# ---- 3. Line (section 22.10) ------------------------------------------------------------------
ws = wb.create_sheet("Line")
ws.append(["orders (x)", "revenue, Rs thousand (y)", "fitted", "residual"])
for c in range(1, 5):
    ws.cell(1, c).font = BOLD
for i, (x, y) in enumerate([(2, 55), (4, 90), (5, 140), (7, 160), (9, 230), (12, 290)], start=2):
    ws.cell(i, 1, x)
    ws.cell(i, 2, y)
    ws.cell(i, 3, f"=$G$2+$G$1*A{i}")
    ws.cell(i, 4, f"=B{i}-C{i}")
ws["A8"], ws["D8"] = "sum of residuals", "=SUM(D2:D7)"
for r, (label, formula) in enumerate([
        ("slope", "=SLOPE(B2:B7,A2:A7)"),
        ("intercept", "=INTERCEPT(B2:B7,A2:A7)"),
        ("R squared", "=RSQ(B2:B7,A2:A7)"),
        ("correlation", "=CORREL(A2:A7,B2:B7)"),
        ("prediction at 10 orders", "=_xlfn.FORECAST.LINEAR(10,B2:B7,A2:A7)")], start=1):
    ws.cell(r, 6, label)
    ws.cell(r, 7, formula)
ws.column_dimensions["B"].width = 24
ws.column_dimensions["F"].width = 24

wb.save(HERE / "ch22_by_hand.xlsx")
print("wrote ch22_by_hand.xlsx")
