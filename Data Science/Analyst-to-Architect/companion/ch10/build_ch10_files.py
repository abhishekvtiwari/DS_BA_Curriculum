"""
Analyst to Architect — Chapter 10: Spreadsheet Fundamentals (Excel & Google Sheets)
build_ch10_files.py — builds every practice file for Chapter 10 from the one-year Riverstone data.

How to run (from this folder):   python3 build_ch10_files.py
Needs: Python 3.10+, openpyxl. Reads ../generate_riverstone_2025.py (seed 20251), so the numbers
match the riverstone_2025 database used in Chapter 13.

Creates:
  riverstone_sales_export_2025.csv   the raw ERP export (one row per order line, 330 rows)
  ch10_practice.xlsx                 Data, Customers, Products, Targets, Cell detective sheets (no formulas)
  ch10_tracker_solution.xlsx         the finished monthly sales tracker from the chapter project
  ch10_tracker_check.xlsx            the same tracker with INDEX/MATCH instead of XLOOKUP (used only to verify
                                     the numbers with LibreOffice, which can't evaluate XLOOKUP)

Tested on: Python 3.12, openpyxl 3.1.5, LibreOffice 24.2 (recalculation of the check file).
Riverstone Supplies is fictional; every name and number is invented.
"""
import csv, datetime as dt, os, runpy, tempfile, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, LineChart, Reference

HERE = pathlib.Path(__file__).resolve().parent
GEN = HERE.parent / "generate_riverstone_2025.py"

# ---- 1. Get the 2025 data from the Chapter 13 generator (run in a temp dir so its SQL file lands there)
cwd = os.getcwd()
with tempfile.TemporaryDirectory() as tmp:
    os.chdir(tmp)
    g = runpy.run_path(str(GEN))
    os.chdir(cwd)

customers = g["customers"]          # (id, name, city, segment, signup, profile, rep, churn)
products = g["products"]            # (id, name, category, price, cost)
employees = {e[0]: e[1] for e in g["employees"]}
orders = {o[0]: o for o in g["new_orders"]}   # (order_id, customer_id, date, status, rep_id)
items = g["items"]                  # (item_id, order_id, product_id, qty, price, disc)
targets = g["targets"]              # (month_start, target)

def code(cid):  # the ERP pads customer IDs to four digits
    return f"{cid:04d}"

# ---- 2. Export rows, in ERP order (order_id, then line)
HEAD = ["order_id", "order_date", "customer_code", "product_id", "quantity", "unit_price",
        "discount_pct", "status", "sales_rep"]
rows = []
for it in items:
    o = orders[it[1]]
    rows.append([o[0], o[2], code(o[1]), it[2], it[3], it[4], it[5], o[3],
                 employees[o[4]] if o[4] else ""])

# CSV exactly as the ERP writes it: DD-MM-YYYY dates, zero-padded codes
with open(HERE / "riverstone_sales_export_2025.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f, lineterminator="\n")
    w.writerow(HEAD)
    for r in rows:
        w.writerow([r[0], r[1].strftime("%d-%m-%Y")] + r[2:])

BOLD = Font(name="Arial", bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="0F5C8C")
ARIAL = Font(name="Arial")

def header(ws, names, row=1):
    for j, n in enumerate(names, start=1):
        c = ws.cell(row=row, column=j, value=n)
        c.font = BOLD; c.fill = HFILL; c.alignment = Alignment(horizontal="center")

def widths(ws, ws_widths):
    for col, wd in ws_widths.items():
        ws.column_dimensions[col].width = wd

def add_reference_sheets(wb):
    ws = wb.create_sheet("Customers")
    header(ws, ["customer_code", "customer_name", "city", "segment", "signup_date"])
    for c in customers:
        ws.append([code(c[0]), c[1], c[2], c[3], c[4]])
    for r in range(2, len(customers) + 2):
        ws.cell(row=r, column=1).number_format = "@"
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
    widths(ws, {"A": 15, "B": 26, "C": 14, "D": 14, "E": 13})
    ws = wb.create_sheet("Products")
    header(ws, ["product_id", "product_name", "category", "list_price", "unit_cost"])
    for p in products:
        ws.append(list(p))
    widths(ws, {"A": 12, "B": 22, "C": 12, "D": 11, "E": 11})
    ws = wb.create_sheet("Targets")
    header(ws, ["month_start", "target_revenue"])
    for m, t in targets:
        ws.append([m, t])
    for r in range(2, 14):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2).number_format = "#,##0"
    widths(ws, {"A": 13, "B": 16})

def export_sheet(wb, title="Export"):
    ws = wb.active; ws.title = title
    header(ws, HEAD)
    for r in rows:
        ws.append(r)
    n = len(rows) + 1
    for r in range(2, n + 1):
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=3).number_format = "@"
    widths(ws, {"A": 10, "B": 12, "C": 14, "D": 11, "E": 10, "F": 11, "G": 13, "H": 11, "I": 16})
    ws.freeze_panes = "A2"
    return ws, n

# ---- 3. Practice workbook (no formulas: the reader adds them)
wb = Workbook()
export_sheet(wb, "Data")
add_reference_sheets(wb)
ws = wb.create_sheet("Cell detective")
header(ws, ["cell_value", "what_it_really_is"])
detective = [
    (2900, "a number"),
    ("2900", "text that looks like a number"),
    ("2900 ", "text with a trailing space"),
    (dt.date(2025, 1, 2), "a real date, shown as a date"),
    ("02-01-2025", "text that looks like a date"),
    (32062.5, "a number with a display format that hides the .5"),
    ("0005", "a code stored as text, so it keeps its zeros"),
]
for v, d in detective:
    ws.append([v, d])
ws["A5"].number_format = "yyyy-mm-dd"
ws["A7"].number_format = "#,##0"
ws["A8"].number_format = "@"
widths(ws, {"A": 16, "B": 46})
wb.save(HERE / "ch10_practice.xlsx")

# ---- 4. Tracker workbook (solution). Two versions: XLOOKUP (delivered) and INDEX/MATCH (for checking)
def build_tracker(path, use_xlookup):
    wb = Workbook()
    ws, n = export_sheet(wb, "Data")
    extra = ["net_revenue", "month_start", "first_line_of_order", "customer_name", "segment"]
    for j, name in enumerate(extra, start=10):
        c = ws.cell(row=1, column=j, value=name); c.font = BOLD; c.fill = HFILL
    for r in range(2, n + 1):
        ws.cell(row=r, column=10, value=f"=E{r}*F{r}*(1-G{r}/100)").number_format = "#,##0.00"
        ws.cell(row=r, column=11, value=f"=DATE(YEAR(B{r}),MONTH(B{r}),1)").number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=12, value=f"=IF(COUNTIF($A$2:A{r},A{r})=1,1,0)")
        if use_xlookup:
            ws.cell(row=r, column=13, value=f'=_xlfn.XLOOKUP(C{r},Customers!$A$2:$A$25,Customers!$B$2:$B$25,"not found")')
            ws.cell(row=r, column=14, value=f'=_xlfn.XLOOKUP(C{r},Customers!$A$2:$A$25,Customers!$D$2:$D$25,"not found")')
        else:
            ws.cell(row=r, column=13, value=f'=IFERROR(INDEX(Customers!$B$2:$B$25,MATCH(C{r},Customers!$A$2:$A$25,0)),"not found")')
            ws.cell(row=r, column=14, value=f'=IFERROR(INDEX(Customers!$D$2:$D$25,MATCH(C{r},Customers!$A$2:$A$25,0)),"not found")')
    widths(ws, {"J": 13, "K": 12, "L": 18, "M": 24, "N": 13})
    add_reference_sheets(wb)
    wb.move_sheet("Customers", offset=0)

    t = wb.create_sheet("Tracker", 0)
    t["A1"] = "Riverstone Supplies — Monthly sales tracker, 2025"; t["A1"].font = Font(name="Arial", bold=True, size=14)
    t["A2"] = "Net revenue = quantity × unit price × (1 − discount %). Cancelled orders excluded. Source: ERP export."
    t["A2"].font = Font(name="Arial", italic=True, color="5B6475")
    cols = ["month", "target", "orders", "net_revenue", "pct_of_target", "vs_prev_month", "status"]
    header(t, cols, row=4)
    last = n
    for i in range(12):
        r = 5 + i
        t.cell(row=r, column=1, value=f"=Targets!A{2+i}").number_format = "mmm yyyy"
        t.cell(row=r, column=2, value=f"=Targets!B{2+i}").number_format = "#,##0"
        t.cell(row=r, column=3, value=(f'=SUMIFS(Data!$L$2:$L${last},Data!$K$2:$K${last},A{r},'
                                       f'Data!$H$2:$H${last},"<>Cancelled")'))
        t.cell(row=r, column=4, value=(f'=SUMIFS(Data!$J$2:$J${last},Data!$K$2:$K${last},A{r},'
                                       f'Data!$H$2:$H${last},"<>Cancelled")')).number_format = "#,##0"
        t.cell(row=r, column=5, value=f"=D{r}/B{r}").number_format = "0.0%"
        if i == 0:
            t.cell(row=r, column=6, value="")
        else:
            t.cell(row=r, column=6, value=f"=D{r}/D{r-1}-1").number_format = "+0.0%;-0.0%"
        t.cell(row=r, column=7, value=f'=IF(D{r}>=B{r},"On target","Below target")')
    t["A17"] = "Total"; t["A17"].font = Font(name="Arial", bold=True)
    for col in "BCD":
        t[f"{col}17"] = f"=SUM({col}5:{col}16)"; t[f"{col}17"].number_format = "#,##0"; t[f"{col}17"].font = Font(name="Arial", bold=True)
    t["E17"] = "=D17/B17"; t["E17"].number_format = "0.0%"; t["E17"].font = Font(name="Arial", bold=True)
    t["A19"] = "Check: total net revenue on the Data sheet (non-cancelled)"
    t["D19"] = f'=SUMIFS(Data!$J$2:$J${last},Data!$H$2:$H${last},"<>Cancelled")'; t["D19"].number_format = "#,##0"
    t["A20"] = "Check: difference (must be 0)"
    t["D20"] = "=D17-D19"; t["D20"].number_format = "#,##0"
    # conditional formatting: green when on target, red when below 80%
    t.conditional_formatting.add("E5:E16", CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                                  fill=PatternFill("solid", fgColor="E2F3EE"), font=Font(color="2F7D6D", bold=True)))
    t.conditional_formatting.add("E5:E16", CellIsRule(operator="lessThan", formula=["0.8"],
                                  fill=PatternFill("solid", fgColor="F8E1E1"), font=Font(color="B23B3B", bold=True)))
    widths(t, {"A": 14, "B": 12, "C": 9, "D": 14, "E": 14, "F": 15, "G": 14})
    t.freeze_panes = "A5"
    # chart: revenue vs target
    ch = BarChart(); ch.type = "col"; ch.title = "Net revenue vs target, 2025 (₹)"
    ch.add_data(Reference(t, min_col=4, min_row=4, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(t, min_col=1, min_row=5, max_row=16))
    ln = LineChart(); ln.add_data(Reference(t, min_col=2, min_row=4, max_row=16), titles_from_data=True)
    ch += ln; ch.height = 8; ch.width = 18
    t.add_chart(ch, "I4")

    # Reps sheet with a data-validated selector
    rp = wb.create_sheet("Reps", 1)
    header(rp, ["sales_rep", "orders", "net_revenue", "share_of_total"], row=1)
    reps = [employees[k] for k in (3, 4, 5)] + [""]
    for i, name in enumerate(reps):
        r = 2 + i
        rp.cell(row=r, column=1, value=name if name else "(no rep)")
        crit = f'A{r}' if name else '""'
        rp.cell(row=r, column=2, value=f'=SUMIFS(Data!$L$2:$L${last},Data!$I$2:$I${last},{crit},Data!$H$2:$H${last},"<>Cancelled")')
        rp.cell(row=r, column=3, value=f'=SUMIFS(Data!$J$2:$J${last},Data!$I$2:$I${last},{crit},Data!$H$2:$H${last},"<>Cancelled")').number_format = "#,##0"
        rp.cell(row=r, column=4, value=f"=C{r}/$C$6").number_format = "0.0%"
    rp["A6"] = "Total"
    rp["B6"] = "=SUM(B2:B5)"; rp["C6"] = "=SUM(C2:C5)"; rp["C6"].number_format = "#,##0"; rp["D6"] = "=SUM(D2:D5)"; rp["D6"].number_format = "0.0%"
    header(rp, ["segment", "orders", "net_revenue", "share_of_total"], row=9)
    for i, seg in enumerate(["Retail", "Hospitality", "Wholesale"]):
        r = 10 + i
        rp.cell(row=r, column=1, value=seg)
        rp.cell(row=r, column=2, value=f'=SUMIFS(Data!$L$2:$L${last},Data!$N$2:$N${last},A{r},Data!$H$2:$H${last},"<>Cancelled")')
        rp.cell(row=r, column=3, value=f'=SUMIFS(Data!$J$2:$J${last},Data!$N$2:$N${last},A{r},Data!$H$2:$H${last},"<>Cancelled")').number_format = "#,##0"
        rp.cell(row=r, column=4, value=f"=C{r}/SUM($C$10:$C$12)").number_format = "0.0%"
    rp["F1"] = "Pick a rep:"; rp["G1"] = employees[3]
    dv = DataValidation(type="list", formula1='"Neha Kulkarni,Rahul Mehta,Farah Khan"', allow_blank=False)
    rp.add_data_validation(dv); dv.add("G1")
    rp["F2"] = "Net revenue:"
    rp["G2"] = f'=SUMIFS(Data!$J$2:$J${last},Data!$I$2:$I${last},G1,Data!$H$2:$H${last},"<>Cancelled")'; rp["G2"].number_format = "#,##0"
    widths(rp, {"A": 16, "B": 9, "C": 14, "D": 15, "F": 13, "G": 16})
    wb.save(path)

build_tracker(HERE / "ch10_tracker_solution.xlsx", use_xlookup=True)
build_tracker(HERE / "ch10_tracker_check.xlsx", use_xlookup=False)
print(f"export rows: {len(rows)}; files written to {HERE}")
