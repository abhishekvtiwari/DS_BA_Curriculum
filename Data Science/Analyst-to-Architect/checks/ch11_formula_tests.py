"""Chapter 11: classic formulas from the text evaluated by LibreOffice on ch11_practice.xlsx.
XLOOKUP, dynamic arrays, QUERY, pivots, Power Query and DAX results are in checks/ch11_expected.py.
Run from the book root:  python3 checks/ch11_formula_tests.py"""
import subprocess, shutil, pathlib
from openpyxl import load_workbook
w = pathlib.Path("/tmp/ch11ft"); shutil.rmtree(w, ignore_errors=True); w.mkdir()
wb = load_workbook("companion/ch11/ch11_practice.xlsx"); T = wb.create_sheet("Tests")
T["E1"] = 502775; T["E2"] = 248727.5; T["E3"] = "Wholesale"; T["E4"] = "Q4"
tests = [
 ("INDEX/MATCH product 105", "=INDEX(Products!B2:B9,MATCH(105,Products!A2:A9,0))"),
 ("INDEX/MATCH city 0014", '=INDEX(Customers!C2:C25,MATCH("0014",Customers!A2:A25,0))'),
 ("approx band 502775", "=INDEX(RebateBands!C2:C4,MATCH(E1,RebateBands!A2:A4,1))"),
 ("approx rate 248727.5", "=INDEX(RebateBands!B2:B4,MATCH(E2,RebateBands!A2:A4,1))"),
 ("VLOOKUP approx", "=VLOOKUP(E1,RebateBands!A2:C4,3,TRUE)"),
 ("SUMIFS segment x quarter", '=SUMIFS(Sales!J2:J331,Sales!N2:N331,E3,Sales!L2:L331,E4,Sales!H2:H331,"<>Cancelled")'),
 ("SUMIFS Wholesale Oct", '=SUMIFS(Sales!J2:J331,Sales!N2:N331,"Wholesale",Sales!K2:K331,DATE(2025,10,1),Sales!H2:H331,"<>Cancelled")'),
 ("valid revenue", '=SUMIFS(Sales!J2:J331,Sales!H2:H331,"<>Cancelled")'),
 ("gross margin", '=SUMIFS(Sales!S2:S331,Sales!H2:H331,"<>Cancelled")'),
 ("last order row Sharma", '=LOOKUP(2,1/(Sales!C2:C331="0001"),Sales!A2:A331)'),
 ("grid price3 vol5", "=4335471*(1+0.03)*(1+0.05)"),
 ("goal seek crates", "=(300000-186928)/(1400*(1-8/100))"),
 ("share wholesale", '=SUMIFS(Sales!J2:J331,Sales!N2:N331,"Wholesale",Sales!H2:H331,"<>Cancelled")/SUMIFS(Sales!J2:J331,Sales!H2:H331,"<>Cancelled")'),
 ("COUNTIFS lines wholesale", '=COUNTIFS(Sales!N2:N331,"Wholesale",Sales!H2:H331,"<>Cancelled")'),
 ("SUMPRODUCT distinct orders wholesale", '=SUMPRODUCT((Sales!N2:N331="Wholesale")*(Sales!H2:H331<>"Cancelled")/COUNTIFS(Sales!A2:A331,Sales!A2:A331))'),
]
for i,(n,f) in enumerate(tests, start=1): T.cell(row=i,column=1,value=n); T.cell(row=i,column=2,value=f)
wb.save(w/"t.xlsx")
subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",str(w/"o"),str(w/"t.xlsx")],capture_output=True)
r = load_workbook(w/"o"/"t.xlsx", data_only=True)["Tests"]
for i,(n,f) in enumerate(tests, start=1): print(f"{n:38} -> {r.cell(row=i,column=2).value!r}")
