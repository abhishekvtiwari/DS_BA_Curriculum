"""Chapter 11 v2 (formula depth): formulas from sections 11.2–11.6 and 11.11 evaluated by LibreOffice on
ch11_practice.xlsx. Functions LibreOffice 24.2 lacks (XMATCH, TEXTSPLIT, TEXTBEFORE/AFTER, LAMBDA family,
TAKE/DROP/CHOOSECOLS) are checked in checks/ch11_expected_v2.py.  Run from the book root."""
import subprocess, shutil, pathlib
from openpyxl import load_workbook
w = pathlib.Path("/tmp/ch11v2"); shutil.rmtree(w, ignore_errors=True); w.mkdir()
wb = load_workbook("companion/ch11/ch11_practice.xlsx"); T = wb.create_sheet("T")
T["H1"] = 50000; T["H2"] = "Wholesale"; T["H3"] = "IN 105 x40 @D6 cost 1102"; T["H4"] = "Deccan Packaging Pvt Ltd"
R = lambda c: f"Sales!{c}2:{c}331"
V = f'{R("H")},"<>Cancelled"'
tests = [
 # 11.2 conditional aggregation
 ("SUMIF wholesale", f'=SUMIF({R("N")},"Wholesale",{R("J")})'),
 ("SUMIF qty>=50 (no sum range)", f'=SUMIF({R("E")},">=50")'),
 ("SUMIFS valid", f'=SUMIFS({R("J")},{V})'),
 ("SUMIFS > cell", f'=SUMIFS({R("J")},{R("J")},">="&H1,{V})'),
 ("SUMIFS between dates Q3", f'=SUMIFS({R("J")},{R("B")},">="&DATE(2025,7,1),{R("B")},"<"&DATE(2025,10,1),{V})'),
 ("SUM(SUMIFS OR array)", f'=SUMPRODUCT(SUMIFS({R("J")},{R("N")},{{"Retail","Hospitality"}},{V}))'),
 ("SUMIFS wildcard *Box*", f'=SUMIFS({R("J")},{R("P")},"*Box*",{V})'),
 ("SUMIFS Storage*", f'=SUMIFS({R("J")},{R("P")},"Storage*",{V})'),
 ("COUNTIFS nonblank rep", f'=COUNTIFS({R("I")},"<>")'),
 ("COUNTIFS blank rep", f'=COUNTIFS({R("I")},"")'),
 ("COUNTIFS between", f'=COUNTIFS({R("J")},">=5000",{R("J")},"<20000",{V})'),
 ("COUNTIFS <5000", f'=COUNTIFS({R("J")},"<5000",{V})'),
 ("COUNTIFS >=20000", f'=COUNTIFS({R("J")},">=20000",{V})'),
 ("AVERAGEIF wholesale line", f'=AVERAGEIF({R("N")},"Wholesale",{R("J")})'),
 ("AVERAGEIFS retail Q4", f'=AVERAGEIFS({R("J")},{R("N")},"Retail",{R("L")},"Q4",{V})'),
 ("MAXIFS Rahul", f'=MAXIFS({R("J")},{R("I")},"Rahul Mehta",{V})'),
 ("MINIFS date Deccan", f'=MINIFS({R("B")},{R("M")},"Deccan Packaging")'),
 ("MAXIFS date Deccan", f'=MAXIFS({R("B")},{R("M")},"Deccan Packaging")'),
 ("SUMPRODUCT month 11", f'=SUMPRODUCT((MONTH({R("B")})=11)*({R("H")}<>"Cancelled")*{R("J")})'),
 ("weighted avg discount", f'=SUMPRODUCT({R("E")}*{R("F")}*{R("G")}*({R("H")}<>"Cancelled"))/SUMPRODUCT({R("E")}*{R("F")}*({R("H")}<>"Cancelled"))'),
 ("simple avg discount", f'=AVERAGEIFS({R("G")},{V})'),
 ("distinct customers", f'=SUMPRODUCT(1/COUNTIF({R("M")},{R("M")}))'),
 ("OR two columns count", f'=SUMPRODUCT(--((({R("G")}>=10)+({R("E")}>=60))>0))'),
 ("gross before discount", f'=SUMPRODUCT({R("E")}*{R("F")}*({R("H")}<>"Cancelled"))'),
 # 11.3 logical, math, stats
 ("SWITCH fiscal", '=SWITCH("Q1","Q1","FY Q4","Q2","FY Q1","Q3","FY Q2","Q4","FY Q3")'),
 ("AGGREGATE 2nd largest retail", f'=AGGREGATE(14,6,{R("J")}/(({R("N")}="Retail")*({R("H")}<>"Cancelled")),2)'),
 ("AGGREGATE largest retail", f'=AGGREGATE(14,6,{R("J")}/(({R("N")}="Retail")*({R("H")}<>"Cancelled")),1)'),
 ("LARGE 2", f'=LARGE({R("J")},2)'), ("LARGE 3", f'=LARGE({R("J")},3)'), ("LARGE 4", f'=LARGE({R("J")},4)'),
 ("SMALL 1", f'=SMALL({R("J")},1)'),
 ("RANK.EQ 51520", f'=RANK.EQ(51520,{R("J")},0)'),
 ("COUNTIF > 51520", f'=COUNTIF({R("J")},">51520")'),
 ("PERCENTILE 0.9", f'=PERCENTILE.INC({R("J")},0.9)'),
 ("QUARTILE 1", f'=QUARTILE.INC({R("J")},1)'), ("QUARTILE 3", f'=QUARTILE.INC({R("J")},3)'),
 ("MROUND", "=MROUND(32062.5,500)"), ("CEILING.MATH", "=CEILING.MATH(32062.5,1000)"), ("FLOOR.MATH", "=FLOOR.MATH(32062.5,1000)"),
 ("QUOTIENT", "=QUOTIENT(45,12)"), ("MOD", "=MOD(45,12)"), ("ROUNDUP cartons", "=ROUNDUP(45/12,0)"),
 ("IFNA", '=IFNA(MATCH("9999",Customers!A2:A25,0),"not found")'),
 ("XOR", "=XOR(TRUE,FALSE)"),
 # 11.4 text and dates
 ("words", '=LEN(TRIM(H4))-LEN(SUBSTITUTE(TRIM(H4)," ",""))+1'),
 ("SEARCH ci", '=SEARCH("pvt",H4)'), ("FIND cs", '=IFERROR(FIND("pvt",H4),"not found")'),
 ("qty from entry", '=VALUE(MID(H3,FIND("x",H3)+1,FIND(" @",H3)-FIND("x",H3)-1))'),
 ("cost from entry", '=VALUE(MID(H3,FIND("cost ",H3)+5,10))'),
 ("REPT bar", '=REPT("|",ROUND(756751/100000,0))'),
 ("EXACT", '=EXACT("Metro Mart","METRO MART")'),
 ("NUMBERVALUE", '=NUMBERVALUE("1.234,50",",",".")'),
 ("TEXTJOIN wholesale", '=TEXTJOIN(", ",TRUE,IF(Customers!D2:D25="Wholesale",Customers!B2:B25,""))'),
 ("EDATE", '=TEXT(EDATE(DATE(2025,1,31),1),"yyyy-mm-dd")'),
 ("WORKDAY", '=TEXT(WORKDAY(DATE(2025,12,24),5),"yyyy-mm-dd")'),
 ("WORKDAY.INTL sun", '=TEXT(WORKDAY.INTL(DATE(2025,12,24),5,11),"yyyy-mm-dd")'),
 ("NETWORKDAYS.INTL sun only", "=NETWORKDAYS.INTL(DATE(2025,12,1),DATE(2025,12,31),11)"),
 ("NETWORKDAYS", "=NETWORKDAYS(DATE(2025,12,1),DATE(2025,12,31))"),
 ("DATEDIF m", "=DATEDIF(DATE(2025,5,6),DATE(2025,12,31),\"m\")"),
 ("DATEDIF md", "=DATEDIF(DATE(2025,5,6),DATE(2025,12,31),\"md\")"),
 ("YEARFRAC", "=ROUND(YEARFRAC(DATE(2025,5,6),DATE(2025,12,31),1),2)"),
 ("WEEKNUM", "=WEEKNUM(DATE(2025,11,6),2)"), ("ISOWEEKNUM", "=ISOWEEKNUM(DATE(2025,11,6))"),
 ("FY end year", "=YEAR(DATE(2025,11,6))+(MONTH(DATE(2025,11,6))>=4)"),
 ("FY label", '="FY"&TEXT(YEAR(DATE(2025,2,10))-(MONTH(DATE(2025,2,10))<4),"0")&"-"&RIGHT(YEAR(DATE(2025,2,10))+(MONTH(DATE(2025,2,10))>=4),2)'),
 ("FY label Nov", '="FY"&TEXT(YEAR(DATE(2025,11,6))-(MONTH(DATE(2025,11,6))<4),"0")&"-"&RIGHT(YEAR(DATE(2025,11,6))+(MONTH(DATE(2025,11,6))>=4),2)'),
 ("FY quarter", "=ROUNDUP(MOD(MONTH(DATE(2025,11,6))-4,12)/3+0.0001,0)"),
 ("FY quarter Jan", "=ROUNDUP(MOD(MONTH(DATE(2025,1,6))-4,12)/3+0.0001,0)"),
 ("FY quarter Apr", "=ROUNDUP(MOD(MONTH(DATE(2025,4,6))-4,12)/3+0.0001,0)"),
 ("days in Feb", "=DAY(EOMONTH(DATE(2025,2,1),0))"),
 ("quarter label", '="Q"&ROUNDUP(MONTH(DATE(2025,11,6))/3,0)'),
 ("customer signup Deccan", '=TEXT(INDEX(Customers!E2:E25,MATCH("0014",Customers!A2:A25,0)),"yyyy-mm-dd")'),
 # 11.5 lookups and references
 ("LOOKUP band", "=LOOKUP(502775,RebateBands!A2:A4,RebateBands!C2:C4)"),
 ("HLOOKUP Oct", '=HLOOKUP("Oct",TargetsWide!B1:M2,2,FALSE)'),
 ("INDEX/MATCH Oct wide", '=INDEX(TargetsWide!B2:M2,MATCH("Oct",TargetsWide!B1:M1,0))'),
 ("OFFSET last 3 targets", "=SUM(OFFSET(Targets!B2,COUNT(Targets!B2:B13)-3,0,3,1))"),
 ("INDEX range last 3", "=SUM(INDEX(Targets!B2:B13,10):INDEX(Targets!B2:B13,12))"),
 ("CHOOSE weekday", '=CHOOSE(WEEKDAY(DATE(2025,11,6),2),"Mon","Tue","Wed","Thu","Fri","Sat","Sun")'),
 ("multi-criteria INDEX/MATCH", f'=INDEX({R("J")},MATCH(1,({R("A")}=10003)*({R("D")}=102),0))'),
 ("SUMIFS as lookup", f'=SUMIFS({R("J")},{R("A")},10003,{R("D")},102)'),
 ("INDEX col 0 sum", "=SUM(INDEX(Customers!A2:E25,0,1)=\"\")"),
 ("LOOKUP last Deccan date", f'=TEXT(LOOKUP(2,1/({R("M")}="Deccan Packaging"),{R("B")}),"yyyy-mm-dd")'),
 ("ADDRESS", "=ADDRESS(5,10)"),
 # 11.11 financial (illustrative)
 ("PMT", "=PMT(10%/12,60,-1200000)"),
 ("total interest", "=PMT(10%/12,60,-1200000)*60-1200000"),
 ("NPV", "=NPV(12%,450000,450000,450000,450000,450000)-1500000"),
 ("IRR", "=IRR({-1500000,450000,450000,450000,450000,450000})"),
 ("FV", "=FV(7%/12,36,-25000)"),
 ("XIRR", "=XIRR({-200000,60000,70000,90000},{45658,45838,46022,46203})"),
]
for i,(n,f) in enumerate(tests, start=1): T.cell(row=i,column=1,value=n); T.cell(row=i,column=2,value=f)
for i,(n,f) in enumerate(tests, start=1):
    if f.startswith(("=TEXTJOIN","=INDEX(Sales","=SUM(INDEX(Customers","=AGGREGATE")) or "1/((" in f or "MATCH(1," in f:
        T.formula_attributes = getattr(T, "formula_attributes", {})
wb.save(w/"t.xlsx")
# LibreOffice: evaluate array formulas by setting them as array formulas in a second pass
from openpyxl.worksheet.formula import ArrayFormula
wb = load_workbook(w/"t.xlsx"); T = wb["T"]
for i,(n,f) in enumerate(tests, start=1):
    if n in ("TEXTJOIN wholesale","multi-criteria INDEX/MATCH","INDEX col 0 sum","AGGREGATE 2nd largest retail","AGGREGATE largest retail","LOOKUP last Deccan date"):
        T.cell(row=i,column=2).value = ArrayFormula(f"B{i}", f)
for fn in ("MAXIFS","MINIFS","SWITCH","IFS","XOR","TEXTJOIN","CEILING.MATH","FLOOR.MATH","ISOWEEKNUM","NUMBERVALUE","AGGREGATE","IFNA","RANK.EQ","PERCENTILE.INC","QUARTILE.INC","WORKDAY.INTL","NETWORKDAYS.INTL"):
    pass
import re
for i,(n,f) in enumerate(tests, start=1):
    c=T.cell(row=i,column=2); v=c.value
    txt = v.text if hasattr(v,"text") else v
    for fn in ["MAXIFS","MINIFS","SWITCH","AGGREGATE","RANK.EQ","PERCENTILE.INC","QUARTILE.INC","CEILING.MATH","FLOOR.MATH","IFNA","XOR","NUMBERVALUE","TEXTJOIN","ISOWEEKNUM","DATEDIF"]:
        if fn=="DATEDIF": continue
        txt = re.sub(r"(?<![\w.])"+re.escape(fn)+r"\(", "_xlfn."+fn+"(", txt)
    if hasattr(v,"text"): v.text = txt
    else: c.value = txt
wb.save(w/"t.xlsx")
subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",str(w/"o"),str(w/"t.xlsx")],capture_output=True)
r = load_workbook(w/"o"/"t.xlsx", data_only=True)["T"]
for i,(n,f) in enumerate(tests, start=1): print(f"{n:36} -> {r.cell(row=i,column=2).value!r}")
