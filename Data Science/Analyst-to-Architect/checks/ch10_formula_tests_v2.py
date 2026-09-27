"""Chapter 10 v2 additions: formulas added in the basics review, evaluated by LibreOffice.
Run from the book root after companion/ch10/build_ch10_files.py:  python3 checks/ch10_formula_tests_v2.py"""
import subprocess, shutil, pathlib
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
work = pathlib.Path("/tmp/ch10v2"); shutil.rmtree(work, ignore_errors=True); work.mkdir()
wb = load_workbook("companion/ch10/ch10_tracker_check.xlsx")
wb.defined_names["net_revenue"] = DefinedName("net_revenue", attr_text="Data!$J$2:$J$331")
T = wb.create_sheet("Tests2")
T["D1"] = "Neha Kulkarni"
tests = [
 ("status bar sum", "=SUM(Data!J2:J331)"), ("status bar avg", "=AVERAGE(Data!J2:J331)"), ("status bar count", "=COUNT(Data!J2:J331)"),
 ("named range sum", "=SUM(net_revenue)"),
 ("MEDIAN all", "=MEDIAN(Data!J2:J331)"),
 ("MAXIFS non-cancelled Retail", '=_xlfn.MAXIFS(Data!J2:J331,Data!N2:N331,"Retail",Data!H2:H331,"<>Cancelled")'),
 ("MINIFS non-cancelled", '=_xlfn.MINIFS(Data!J2:J331,Data!H2:H331,"<>Cancelled")'),
 ("ROUNDUP -2", "=ROUNDUP(Data!J5,-2)"), ("ROUNDDOWN -2", "=ROUNDDOWN(Data!J5,-2)"), ("INT", "=INT(Data!J5)"),
 ("DIV0", "=Data!J2/0"), ("VALUE err", '="abc"*2'), ("NAME err", "=SUMM(Data!J2:J5)"),
 ("NA err", '=INDEX(Customers!B2:B25,MATCH("9999",Customers!A2:A25,0))'),
 ("TIME value", "=TIME(9,30,0)"), ("hours worked", "=(TIME(17,30,0)-TIME(9,15,0))*24"),
 ("first name", '=LEFT(D1,FIND(" ",D1)-1)'), ("last name", '=MID(D1,FIND(" ",D1)+1,100)'),
 ("CONCAT", '=_xlfn.CONCAT("RS-",Data!A2)'),
 ("SUM with insert-safe range", "=SUM(Data!J2:J331)"),
]
for i,(n,f) in enumerate(tests, start=1):
    T.cell(row=i, column=1, value=n); T.cell(row=i, column=2, value=f)
wb.save(work/"t.xlsx")
subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",str(work/"out"),str(work/"t.xlsx")],capture_output=True)
r = load_workbook(work/"out"/"t.xlsx", data_only=True)["Tests2"]
for i,(n,f) in enumerate(tests, start=1):
    print(f"{n:30} {f:55.55} -> {r.cell(row=i,column=2).value!r}")
