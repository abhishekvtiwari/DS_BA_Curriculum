"""Chapter 15: check the spreadsheet formulas the chapter quotes, in LibreOffice (headless), not by guessing.

Run from the book folder:  python3 checks/ch15_spreadsheet_check.py
Needs: pandas, openpyxl, LibreOffice (soffice). openpyxl writes QUARTILE.INC as _xlfn.QUARTILE.INC (the file-format name). LibreOffice 24.2 has no FILTER, so each segment's order values
are written to their own column; QUARTILE.INC/MEDIAN on that column is what FILTER(...) hands them in Excel/Sheets.
"""
import pathlib, subprocess, tempfile, csv
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L

BOOK = pathlib.Path(__file__).resolve().parents[1]
D = BOOK / "companion" / "ch15" / "chart_data"
ov = pd.read_csv(D / "order_values_2025.csv"); cu = pd.read_csv(D / "customers_2025.csv")
mon = pd.read_csv(D / "monthly_2023_2025.csv"); ans = pd.read_csv(D / "anscombe.csv")

wb = Workbook(); ws = wb.active; ws.title = "check"; data = wb.create_sheet("data")
palm = [9678, 15136, 19292, 20520, 25865, 27960, 39975, 62335, 99680]      # Palm Trading Co, 2025 (section 15.5)
cols = {"palm": palm}
for s in ["Hospitality", "Retail", "Wholesale"]:
    cols[s] = ov[ov.segment == s].order_value.tolist()
cols["orders"] = cu.orders.tolist(); cols["cust_rev"] = cu.net_revenue.tolist()
m25 = mon[mon.year == 2025]; cols["rev25"] = m25.net_revenue.tolist(); cols["tgt25"] = m25.target_revenue.tolist()
for k in ["x_1_2_3", "y1", "y2", "y3", "x4", "y4"]: cols[k] = ans[k].tolist()
ref = {}
for j, (k, vals) in enumerate(cols.items(), 1):
    data.cell(1, j, k)
    for i, v in enumerate(vals, 2): data.cell(i, j, v)
    ref[k] = f"data!{L(j)}2:{L(j)}{len(vals) + 1}"
checks = [("palm median", f"=MEDIAN({ref['palm']})"), ("palm Q1", f"=_xlfn.QUARTILE.INC({ref['palm']},1)"),
          ("palm Q3", f"=_xlfn.QUARTILE.INC({ref['palm']},3)")]
for s in ["Hospitality", "Retail", "Wholesale"]:
    checks += [(f"{s} Q1", f"=_xlfn.QUARTILE.INC({ref[s]},1)"), (f"{s} median", f"=MEDIAN({ref[s]})"),
               (f"{s} Q3", f"=_xlfn.QUARTILE.INC({ref[s]},3)"), (f"{s} max", f"=MAX({ref[s]})")]
checks += [("CORREL orders vs revenue", f"=CORREL({ref['orders']},{ref['cust_rev']})"),
           ("2025 attainment (B14)", f"=SUM({ref['rev25']})/SUM({ref['tgt25']})"),
           ("title (C14)", f'="2025 finished at "&TEXT(SUM({ref["rev25"]})/SUM({ref["tgt25"]}),"0.0%")&" of target"')]
for k in ["y1", "y2", "y3"]:
    checks += [(f"Anscombe mean {k}", f"=AVERAGE({ref[k]})"), (f"Anscombe CORREL x,{k}", f"=CORREL({ref['x_1_2_3']},{ref[k]})")]
checks += [("Anscombe mean y4", f"=AVERAGE({ref['y4']})"), ("Anscombe CORREL x4,y4", f"=CORREL({ref['x4']},{ref['y4']})")]
for i, (name, f) in enumerate(checks, 1):
    ws.cell(i, 1, name); ws.cell(i, 2, f)
with tempfile.TemporaryDirectory() as t:
    x = pathlib.Path(t) / "ch15_check.xlsx"; wb.save(x)
    subprocess.run(["soffice", "--headless", "--convert-to", "csv", "--outdir", t, str(x)], check=True, capture_output=True)
    for row in csv.reader(open(pathlib.Path(t) / "ch15_check.csv", encoding="utf-8")):
        print(f"{row[0]:28} {row[1]}")
