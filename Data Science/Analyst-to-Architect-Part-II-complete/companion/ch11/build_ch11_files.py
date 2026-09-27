"""
Analyst to Architect — Chapter 11: The Spreadsheet, Mastered (Excel & Google Sheets)
build_ch11_files.py — builds every practice file for Chapter 11 from the one-year Riverstone data.

How to run (from this folder):   python3 build_ch11_files.py
Needs: Python 3.10+, openpyxl. Reads ../generate_riverstone_2025.py (seed 20251), so every total matches
Chapter 10's tracker and the riverstone_2025 database in Chapter 13.

Creates:
  monthly_exports/sales_2025_01.csv … sales_2025_12.csv   one ERP export per month (the Power Query folder)
  ch11_practice.xlsx              clean Sales table plus Customers, Products, Targets, TargetsWide, RebateBands
  month_end_pack_2025_messy.xlsx  the copy-paste month-end workbook from the chapter story (planted errors)
  riverstone_sales.pq             the Power Query (M) code shown in section 11.7
  riverstone_measures.dax         the DAX measures shown in section 11.8

Tested on: Python 3.12, openpyxl 3.1.5; the messy workbook's SUM formulas recalculated in LibreOffice 24.2.
Riverstone Supplies is fictional; every name and number is invented.
"""
import csv, os, runpy, tempfile, pathlib, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = pathlib.Path(__file__).resolve().parent
GEN = HERE.parent / "generate_riverstone_2025.py"
cwd = os.getcwd()
with tempfile.TemporaryDirectory() as tmp:
    os.chdir(tmp); g = runpy.run_path(str(GEN)); os.chdir(cwd)

customers = g["customers"]; products = g["products"]
employees = {e[0]: e[1] for e in g["employees"]}
orders = {o[0]: o for o in g["new_orders"]}; items = g["items"]; targets = g["targets"]
cust = {c[0]: c for c in customers}; prod = {p[0]: p for p in products}
code = lambda cid: f"{cid:04d}"
HEAD = ["order_id", "order_date", "customer_code", "product_id", "quantity", "unit_price", "discount_pct", "status", "sales_rep"]

rows = []
for it in items:
    o = orders[it[1]]
    rows.append([o[0], o[2], code(o[1]), it[2], it[3], it[4], it[5], o[3], employees[o[4]] if o[4] else ""])

# ---- 1. monthly CSV exports
out = HERE / "monthly_exports"; out.mkdir(exist_ok=True)
for m in range(1, 13):
    with open(out / f"sales_2025_{m:02d}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(HEAD)
        for r in rows:
            if r[1].month == m:
                w.writerow([r[0], r[1].strftime("%d-%m-%Y")] + r[2:])

BOLD = Font(name="Arial", bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="0F5C8C")
def header(ws, names, row=1):
    for j, n in enumerate(names, start=1):
        c = ws.cell(row=row, column=j, value=n); c.font = BOLD; c.fill = HFILL
def net(r): return r[4] * r[5] * (1 - r[6] / 100)

# ---- 2. practice workbook: one clean, wide Sales table (values, not formulas) plus lookups
wb = Workbook(); ws = wb.active; ws.title = "Sales"
cols = HEAD + ["net_revenue", "month_start", "quarter", "customer_name", "segment", "city", "product_name", "category", "unit_cost", "gross_margin"]
header(ws, cols)
for r in rows:
    c = cust[int(r[2])]; p = prod[r[3]]
    nr = round(net(r), 2)
    ws.append(r + [nr, dt.date(r[1].year, r[1].month, 1), f"Q{(r[1].month - 1)//3 + 1}", c[1], c[3], c[2],
                   p[1], p[2], p[4], round(nr - r[4] * p[4], 2)])
for rr in range(2, len(rows) + 2):
    ws.cell(row=rr, column=2).number_format = "yyyy-mm-dd"; ws.cell(row=rr, column=11).number_format = "yyyy-mm-dd"
    ws.cell(row=rr, column=3).number_format = "@"
ws.freeze_panes = "A2"
cs = wb.create_sheet("Customers"); header(cs, ["customer_code", "customer_name", "city", "segment", "signup_date"])
for c in customers: cs.append([code(c[0]), c[1], c[2], c[3], c[4]])
for rr in range(2, 26): cs.cell(row=rr, column=1).number_format = "@"; cs.cell(row=rr, column=5).number_format = "yyyy-mm-dd"
ps = wb.create_sheet("Products"); header(ps, ["product_id", "product_name", "category", "list_price", "unit_cost"])
for p in products: ps.append(list(p))
ts = wb.create_sheet("Targets"); header(ts, ["month_start", "target_revenue"])
for m, t in targets: ts.append([m, t])
for rr in range(2, 14): ts.cell(row=rr, column=1).number_format = "yyyy-mm-dd"
tw = wb.create_sheet("TargetsWide"); header(tw, ["year"] + [dt.date(2025, m, 1).strftime("%b") for m in range(1, 13)])
tw.append([2025] + [t for _, t in targets])
rb = wb.create_sheet("RebateBands"); header(rb, ["annual_revenue_from", "rebate_pct", "band"])
for a, b, c in [(0, 0, "Standard"), (250000, 1, "Silver"), (400000, 2, "Gold")]: rb.append([a, b, c])
wb.save(HERE / "ch11_practice.xlsx")

# ---- 3. the messy month-end workbook (one sheet per month, pasted values, a hand-built Summary)
MONTHS = [dt.date(2025, m, 1).strftime("%b") for m in range(1, 13)]
mw = Workbook(); summ = mw.active; summ.title = "Summary"
sheet_last = {}
for m in range(1, 13):
    sh = mw.create_sheet(MONTHS[m - 1]); header(sh, HEAD + ["net_revenue"])
    mrows = [r for r in rows if r[1].month == m]
    if m == 12:                     # planted: the last five December lines were pasted twice
        mrows = mrows + mrows[-5:]
    for i, r in enumerate(mrows, start=2):
        sh.append(r); sh.cell(row=i, column=2).number_format = "dd-mm-yyyy"
        sh.cell(row=i, column=10, value=f"=E{i}*F{i}*(1-G{i}/100)")
    sheet_last[m] = len(mrows) + 1
summ["A1"] = "Riverstone month-end pack 2025 (built by copy and paste)"; summ["A1"].font = Font(name="Arial", bold=True, size=13)
header(summ, ["month", "net_revenue", "target", "pct_of_target", "note"], row=3)
for m in range(1, 13):
    r = 3 + m; last = sheet_last[m]; name = MONTHS[m - 1]
    summ.cell(row=r, column=1, value=name)
    if m in (1, 2, 5, 6, 7, 8, 9, 12):
        f = f'=SUMIFS({name}!J2:J{last},{name}!H2:H{last},"<>Cancelled")'
    elif m in (3, 4):               # planted: March and April use SUM, so cancelled orders are included
        f = f"=SUM({name}!J2:J{last})"
    elif m == 10:                   # planted: October's range stops at row 30
        f = f'=SUMIFS({name}!J2:J30,{name}!H2:H30,"<>Cancelled")'
    elif m == 11:                   # planted: November typed in by hand, digits transposed
        f = 633480
    summ.cell(row=r, column=2, value=f).number_format = "#,##0"
    summ.cell(row=r, column=3, value=targets[m - 1][1]).number_format = "#,##0"
    summ.cell(row=r, column=4, value=f"=B{r}/C{r}").number_format = "0.0%"
summ["E14"] = "typed from the flash email"
summ["A16"] = "Total"; summ["B16"] = "=SUM(B4:B15)"; summ["C16"] = "=SUM(C4:C15)"; summ["D16"] = "=B16/C16"
summ["B16"].number_format = "#,##0"; summ["C16"].number_format = "#,##0"; summ["D16"].number_format = "0.0%"
mw.save(HERE / "month_end_pack_2025_messy.xlsx")

# ---- 4. Power Query and DAX text files (documentation; run them in Excel)
(HERE / "riverstone_sales.pq").write_text('''// Analyst to Architect, Chapter 11, section 11.7 — Power Query (M) for Riverstone's monthly exports.
// Paste into Excel: Data > Get Data > From Other Sources > Blank Query > Advanced Editor. Change FolderPath.
// Riverstone Supplies is fictional; every name and number is invented.
let
    FolderPath = "C:\\Riverstone\\monthly_exports",
    Source     = Folder.Files(FolderPath),
    CsvOnly    = Table.SelectRows(Source,
                     each Text.Lower([Extension]) = ".csv"),
    Parsed     = Table.AddColumn(CsvOnly, "Data",
                     each Table.PromoteHeaders(
                         Csv.Document([Content],
                             [Delimiter = ",", Encoding = 65001,
                              QuoteStyle = QuoteStyle.Csv]),
                         [PromoteAllScalars = true])),
    Combined   = Table.Combine(Parsed[Data]),
    Typed      = Table.TransformColumnTypes(Combined, {
                     {"order_id", Int64.Type},
                     {"customer_code", type text},
                     {"product_id", Int64.Type},
                     {"quantity", Int64.Type},
                     {"unit_price", type number},
                     {"discount_pct", type number},
                     {"status", type text},
                     {"sales_rep", type text}}),
    TypedDates = Table.TransformColumnTypes(Typed,
                     {{"order_date", type date}}, "en-IN"),
    Valid      = Table.SelectRows(TypedDates,
                     each [status] <> "Cancelled"),
    NetRevenue = Table.AddColumn(Valid, "net_revenue",
                     each [quantity] * [unit_price]
                          * (1 - [discount_pct] / 100),
                     type number)
in
    NetRevenue
''', encoding="utf-8")
(HERE / "riverstone_measures.dax").write_text('''// Analyst to Architect, Chapter 11, section 11.8 — DAX measures for the Riverstone data model.
// Tables: Sales (one row per order line), Customers, Products, Calendar (one row per date in 2025).
// Relationships: Sales[customer_code] -> Customers[customer_code]; Sales[product_id] -> Products[product_id];
//                Sales[order_date] -> Calendar[Date] (Calendar marked as date table).
// Riverstone Supplies is fictional; every name and number is invented.

Net Revenue :=
SUMX ( Sales, Sales[quantity] * Sales[unit_price] * ( 1 - Sales[discount_pct] / 100 ) )

Valid Net Revenue :=
CALCULATE ( [Net Revenue], Sales[status] <> "Cancelled" )

Gross Margin :=
CALCULATE (
    SUMX ( Sales,
        Sales[quantity] * ( Sales[unit_price] * ( 1 - Sales[discount_pct] / 100 ) - RELATED ( Products[unit_cost] ) ) ),
    Sales[status] <> "Cancelled" )

Gross Margin % :=
DIVIDE ( [Gross Margin], [Valid Net Revenue] )

Wholesale Revenue :=
CALCULATE ( [Valid Net Revenue], Customers[segment] = "Wholesale" )

Revenue YTD :=
TOTALYTD ( [Valid Net Revenue], Calendar[Date] )
''', encoding="utf-8")
print(f"lines: {len(rows)}; monthly files: 12; written to {HERE}")

# ---- 5. Championship-style case "The Riverstone Stockroom" (section 11.14): case file and a helper-column solution model
CH = HERE / "challenge"; CH.mkdir(exist_ok=True)
STATUS = {"Delivered": "DEL", "Shipped": "SHP", "Pending": "PND", "Cancelled": "CAN"}
def logcode(r):
    return f"{r[1].strftime('%d.%m.%Y')}|SO{r[0]}|C{r[2]}|P{r[3]}x{r[4]}@{int(r[5])}-{int(r[6])}%|{STATUS[r[7]]}"
case = Workbook(); cs = case.active; cs.title = "Case"
brief = [
 "THE RIVERSTONE STOCKROOM — a championship-style case (30 minutes, 1,000 points + bonus)",
 "",
 "Riverstone's warehouse wants to know whether a simple weekly reorder rule would have kept Storage Box 10L (product 101) in stock in 2025.",
 "The Log sheet has every 2025 order line as one text code: date|order|customer|product x quantity @ price - discount %|status",
 "Example: 02.01.2025|SO10001|C0002|P108x10@290-0%|DEL  (status codes: DEL delivered, SHP shipped, PND pending, CAN cancelled)",
 "",
 "RULES FOR LEVELS 2 TO 5 (product 101 only; cancelled lines never ship)",
 "1. The model runs day by day from 1 Jan 2025 to 31 Dec 2025. Demand on a day = units of product 101 on non-cancelled lines dated that day.",
 "2. Opening stock on 1 Jan 2025 is 300 units. Closing stock = opening stock + units received that day - demand. Stock may go negative (a backorder).",
 "3. Reorder rule: every Monday, if the previous day's closing stock is less than or equal to the reorder point (ROP), a purchase order (PO) for Q units is placed.",
 "4. A PO placed on day d arrives at the start of day d + L (L = lead time in days) and counts as received that day.",
 "5. Costs: holding Rs 1 per unit per day of positive closing stock; shortage Rs 20 per unit per day of negative closing stock; Rs 2,000 per PO.",
 "",
 "LEVEL 1 (100 points)  a) Total units on non-cancelled lines, all products.  b) Product ID with the most units.  c) For product 101, the customer code that bought the most units.",
 "LEVEL 2 (150 points)  With no reordering at all, on which date does product 101's closing stock first go below zero? What is the closing stock on 31 Dec?",
 "LEVEL 3 (200 points)  With ROP = 100, Q = 250, L = 7: how many POs are placed, and what is the lowest closing stock and its first date?",
 "LEVEL 4 (250 points)  With ROP = 100, Q = 250, L = 7: on how many days is closing stock negative? Keeping Q = 250, what is the smallest ROP (multiple of 50) with no negative day?",
 "LEVEL 5 (300 points)  L = 7. ROP can be 0, 50, ..., 400 and Q can be 50, 100, ..., 800. Which combination gives the lowest total cost for 2025, and what is that cost?",
 "BONUS (50 points)     Same search as Level 5 with L = 14. What is the lowest total cost?",
]
for line in brief: cs.append([line])
cs.column_dimensions["A"].width = 160
lg = case.create_sheet("Log"); lg.append(["code"])
for r in rows: lg.append([logcode(r)])
lg.column_dimensions["A"].width = 48
pa = case.create_sheet("Parameters")
for row in [["opening_stock", 300], ["product_id", 101], ["rop", 100], ["q", 250], ["lead_days", 7],
            ["holding_per_unit_day", 1], ["shortage_per_unit_day", 20], ["cost_per_po", 2000]]:
    pa.append(row)
an = case.create_sheet("Answers"); an.append(["level", "question", "your answer"])
for lvl, qn in [(1, "a"), (1, "b"), (1, "c"), (2, "first negative date"), (2, "31 Dec closing"), (3, "POs"), (3, "lowest closing"), (3, "date of lowest"),
                (4, "negative days"), (4, "smallest ROP"), (5, "ROP"), (5, "Q"), (5, "cost"), ("bonus", "cost")]:
    an.append([lvl, qn, None])
case.save(CH / "riverstone_stockroom_case.xlsx")

# helper-column solution (classic functions only, so it runs in every Excel version, Sheets and LibreOffice)
sol = Workbook(); L = sol.active; L.title = "Log"
L.append(["code", "order_date", "order_id", "customer_code", "product_id", "quantity", "unit_price", "discount_pct", "status"])
n = len(rows) + 1
for i, r in enumerate(rows, start=2):
    L.cell(row=i, column=1, value=logcode(r))
    L.cell(row=i, column=2, value=f"=DATE(MID(A{i},7,4),MID(A{i},4,2),LEFT(A{i},2))").number_format = "yyyy-mm-dd"
    L.cell(row=i, column=3, value=f"=VALUE(MID(A{i},14,5))")
    L.cell(row=i, column=4, value=f"=MID(A{i},21,4)")
    L.cell(row=i, column=5, value=f"=VALUE(MID(A{i},27,3))")
    L.cell(row=i, column=6, value=f'=VALUE(MID(A{i},31,FIND("@",A{i})-31))')
    L.cell(row=i, column=7, value=f'=VALUE(MID(A{i},FIND("@",A{i})+1,FIND("-",A{i})-FIND("@",A{i})-1))')
    L.cell(row=i, column=8, value=f'=VALUE(MID(A{i},FIND("-",A{i})+1,FIND("%",A{i})-FIND("-",A{i})-1))')
    L.cell(row=i, column=9, value=f"=RIGHT(A{i},3)")
P = sol.create_sheet("Parameters")
for row in [["opening_stock", 300], ["product_id", 101], ["rop", 100], ["q", 250], ["lead_days", 7],
            ["holding_per_unit_day", 1], ["shortage_per_unit_day", 20], ["cost_per_po", 2000]]:
    P.append(row)
Mo = sol.create_sheet("Model")
Mo.append(["date", "weekday", "demand", "receipts", "po_placed", "closing"])
for d in range(365):
    r = d + 2
    Mo.cell(row=r, column=1, value=(f"=DATE(2025,1,1)" if d == 0 else f"=A{r-1}+1")).number_format = "yyyy-mm-dd"
    Mo.cell(row=r, column=2, value=f"=WEEKDAY(A{r},2)")
    Mo.cell(row=r, column=3, value=f'=SUMIFS(Log!$F$2:$F${n},Log!$B$2:$B${n},A{r},Log!$E$2:$E${n},Parameters!$B$2,Log!$I$2:$I${n},"<>CAN")')
    Mo.cell(row=r, column=4, value=f"=IF(ROW()-Parameters!$B$5>=2,INDEX($E$1:$E$366,ROW()-Parameters!$B$5)*Parameters!$B$4,0)")
    Mo.cell(row=r, column=5, value=("=0" if d == 0 else f"=IF(AND(B{r}=1,F{r-1}<=Parameters!$B$3),1,0)"))
    Mo.cell(row=r, column=6, value=(f"=Parameters!$B$1+D{r}-C{r}" if d == 0 else f"=F{r-1}+D{r}-C{r}"))
Sm = sol.create_sheet("Results")
res = [
 ("L1a total units, non-cancelled", f'=SUMIFS(Log!F2:F{n},Log!I2:I{n},"<>CAN")'),
 ("L1b units of product 101 (compare each product)", f'=SUMIFS(Log!F2:F{n},Log!E2:E{n},101,Log!I2:I{n},"<>CAN")'),
 ("L1c units of 101 bought by C0012", f'=SUMIFS(Log!F2:F{n},Log!E2:E{n},101,Log!D2:D{n},"0012",Log!I2:I{n},"<>CAN")'),
 ("POs placed", "=SUM(Model!E2:E366)"),
 ("lowest closing", "=MIN(Model!F2:F366)"),
 ("date of lowest", "=INDEX(Model!A2:A366,MATCH(MIN(Model!F2:F366),Model!F2:F366,0))"),
 ("negative days", '=COUNTIF(Model!F2:F366,"<0")'),
 ("31 Dec closing", "=Model!F366"),
 ("holding cost", '=SUMIF(Model!F2:F366,">0")*Parameters!B6'),
 ("shortage cost", '=-SUMIF(Model!F2:F366,"<0")*Parameters!B7'),
 ("ordering cost", "=B4*Parameters!B8"),
 ("total cost", "=B9+B10+B11"),
]
for i, (k, f) in enumerate(res, start=1):
    Sm.cell(row=i, column=1, value=k); Sm.cell(row=i, column=2, value=f)
sol.save(CH / "riverstone_stockroom_solution.xlsx")
print("challenge files written")

# ---- 6. Power Query practice: an export with problems (section 11.7, "When a file breaks the rules")
PQ = HERE / "pq_practice"; PQ.mkdir(exist_ok=True)
dec_tail = [r for r in rows if r[1].month == 12][-5:]
with open(PQ / "sales_export_with_problems.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(HEAD + ["remarks"])
    for i, r in enumerate(dec_tail):
        qty = f"{r[4]} pcs" if i == 2 else r[4]
        w.writerow([r[0], r[1].strftime("%d-%m-%Y"), r[2], r[3], qty, r[5], r[6], r[7], r[8], "urgent" if i == 0 else ""])
    w.writerow(["Total lines: 5", "", "", "", "", "", "", "", "", ""])
print("pq practice file written")

# ---- 7. Timed practice case: Riverstone Rewards (section 11.14) — case file and a classic-formula solution
from openpyxl.styles import Font as _F
RULES = [
    "RIVERSTONE REWARDS — competition-style case (Analyst to Architect, Chapter 11, section 11.14). Target time: 45 minutes.",
    "Data: the Lines sheet has all 330 order lines of 2025. Cancelled orders earn nothing and don't exist for any rule.",
    "Rule 1  Order value = sum of net_revenue over the order's lines.",
    "Rule 2  Base points = order value / 100, rounded down to a whole number.",
    "Rule 3  Tier at the time of an order depends on the customer's total order value from their EARLIER orders in 2025",
    "        (orders with a smaller order_id): below 100,000 = Standard (x1); 100,000 or more = Silver (x1.25);",
    "        250,000 or more = Gold (x1.5). Tier points = base points x multiplier, rounded down.",
    "Rule 4  Streak bonus: +200 points on an order if the same customer placed at least one order in the previous calendar month.",
    "Rule 5  Basket bonus: +100 points on an order with 3 or more lines.",
    "Rule 6  Order points = tier points + streak bonus + basket bonus.",
    "Rule 7  Expiry on 31 Dec 2025: points from orders dated before 1 Jul 2025 expire, unless the customer ordered on or after 1 Oct 2025.",
    "Rule 8  Balance = all order points minus expired points.",
]
QUESTIONS = [
    ("L1", "How many orders earn points, and how many of them are worth 25,000 or more?"),
    ("L2", "Total base points across all orders? Which customer has the most base points?"),
    ("L3", "How many orders were placed at Gold tier? At Silver?"),
    ("L4", "Total tier points (after multipliers, before bonuses)?"),
    ("L5", "How many orders earn the streak bonus? Total points after tier and streak (no basket bonus yet)?"),
    ("L6", "Total order points (Rule 6) across all orders?"),
    ("L7", "How many customers lose points to expiry, and how many points expire in total?"),
    ("L8", "Highest balance (customer and points)? How many customers finish with 4,000 points or more?"),
    ("B1", "Bonus: which customer's running points first reach 5,000, on which order_id?"),
    ("B2", "Bonus: which month earned the most order points, and how many?"),
    ("B3", "Bonus: how many customers ordered in every month of 2025?"),
    ("B4", "Bonus: Metro Mart's rank by balance (1 = highest)?"),
]
cw = Workbook(); r0 = cw.active; r0.title = "Rules"
for i, t in enumerate(RULES, start=1):
    r0.cell(row=i, column=1, value=t).font = _F(name="Arial", bold=(i == 1))
q0 = cw.create_sheet("Questions"); header(q0, ["level", "question", "your_answer"])
for lv, q in QUESTIONS: q0.append([lv, q, None])
q0.column_dimensions["B"].width = 110; q0.column_dimensions["C"].width = 24
ln = cw.create_sheet("Lines"); header(ln, ["order_id", "order_date", "customer_name", "product_id", "quantity", "unit_price", "discount_pct", "status", "net_revenue"])
for r in rows:
    ln.append([r[0], r[1], cust[int(r[2])][1], r[3], r[4], r[5], r[6], r[7], round(net(r), 2)])
for rr in range(2, len(rows) + 2): ln.cell(row=rr, column=2).number_format = "yyyy-mm-dd"
cw.save(CH / "riverstone_rewards_case.xlsx")

# Solution with classic formulas only (works in Excel 2010+, Google Sheets, LibreOffice)
sw = Workbook(); L = sw.active; L.title = "Lines"
for row in ln.iter_rows(values_only=True): L.append(row)
for rr in range(2, len(rows) + 2): L.cell(row=rr, column=2).number_format = "yyyy-mm-dd"
N = len(rows) + 1
valid_ids = sorted({r[0] for r in rows if r[7] != "Cancelled"})
Ow = sw.create_sheet("Orders")
header(Ow, ["order_id", "order_date", "customer", "value", "lines", "month", "base", "prev_spend", "multiplier",
            "tier_points", "streak", "basket", "points", "running_points", "expired", "hit_5000"])
last = len(valid_ids) + 1
for i, oid in enumerate(valid_ids, start=2):
    Ow.cell(row=i, column=1, value=oid)
    f = {
        2: f"=INDEX(Lines!$B$2:$B${N},MATCH(A{i},Lines!$A$2:$A${N},0))",
        3: f"=INDEX(Lines!$C$2:$C${N},MATCH(A{i},Lines!$A$2:$A${N},0))",
        4: f"=SUMIFS(Lines!$I$2:$I${N},Lines!$A$2:$A${N},A{i})",
        5: f"=COUNTIFS(Lines!$A$2:$A${N},A{i})",
        6: f"=DATE(YEAR(B{i}),MONTH(B{i}),1)",
        7: f"=INT(D{i}/100)",
        8: f"=SUMIFS($D$2:D{i},$C$2:C{i},C{i})-D{i}",
        9: f"=IF(H{i}>=250000,1.5,IF(H{i}>=100000,1.25,1))",
        10: f"=ROUNDDOWN(G{i}*I{i},0)",
        11: f"=IF(COUNTIFS($C$2:$C${last},C{i},$F$2:$F${last},EDATE(F{i},-1))>0,200,0)",
        12: f"=IF(E{i}>=3,100,0)",
        13: f"=J{i}+K{i}+L{i}",
        14: f"=SUMIFS($M$2:M{i},$C$2:C{i},C{i})",
        15: f'=IF(AND(B{i}<DATE(2025,7,1),COUNTIFS($C$2:$C${last},C{i},$B$2:$B${last},">="&DATE(2025,10,1))=0),M{i},0)',
        16: f'=IF(N{i}>=5000,A{i},"")',
    }
    for col, fx in f.items(): Ow.cell(row=i, column=col, value=fx)
    Ow.cell(row=i, column=2).number_format = "yyyy-mm-dd"; Ow.cell(row=i, column=6).number_format = "yyyy-mm-dd"
Cs = sw.create_sheet("Customers"); header(Cs, ["customer", "points", "expired", "balance", "rank", "months_ordered"])
names = sorted({cust[int(r[2])][1] for r in rows if r[7] != "Cancelled"})
cl = len(names) + 1
for i, nm in enumerate(names, start=2):
    Cs.cell(row=i, column=1, value=nm)
    Cs.cell(row=i, column=2, value=f"=SUMIFS(Orders!$M$2:$M${last},Orders!$C$2:$C${last},A{i})")
    Cs.cell(row=i, column=3, value=f"=SUMIFS(Orders!$O$2:$O${last},Orders!$C$2:$C${last},A{i})")
    Cs.cell(row=i, column=4, value=f"=B{i}-C{i}")
    Cs.cell(row=i, column=5, value=f"=RANK(D{i},$D$2:$D${cl},0)")
    Cs.cell(row=i, column=6, value=f"=SUMPRODUCT((Orders!$C$2:$C${last}=A{i})/COUNTIFS(Orders!$C$2:$C${last},Orders!$C$2:$C${last},Orders!$F$2:$F${last},Orders!$F$2:$F${last}))")
Ms = sw.create_sheet("Months"); header(Ms, ["month", "points"])
for m in range(1, 13):
    Ms.cell(row=m + 1, column=1, value=dt.date(2025, m, 1)).number_format = "yyyy-mm"
    Ms.cell(row=m + 1, column=2, value=f"=SUMIFS(Orders!$M$2:$M${last},Orders!$F$2:$F${last},A{m+1})")
An = sw.create_sheet("Answers", 0); header(An, ["level", "item", "answer"])
ans = [
    ("L1", "orders", f"=COUNT(Orders!A2:A{last})"), ("L1", "orders >= 25,000", f'=COUNTIF(Orders!D2:D{last},">=25000")'),
    ("L2", "total base points", f"=SUM(Orders!G2:G{last})"),
    ("L2", "top customer by base points", f"=INDEX(Customers!A2:A{cl},MATCH(MAX(Customers!G2:G{cl}),Customers!G2:G{cl},0))"),
    ("L3", "Gold orders", f'=COUNTIF(Orders!I2:I{last},1.5)'), ("L3", "Silver orders", f'=COUNTIF(Orders!I2:I{last},1.25)'),
    ("L4", "tier points", f"=SUM(Orders!J2:J{last})"),
    ("L5", "streak orders", f'=COUNTIF(Orders!K2:K{last},200)'), ("L5", "tier + streak points", f"=SUM(Orders!J2:K{last})"),
    ("L6", "total order points", f"=SUM(Orders!M2:M{last})"),
    ("L7", "customers losing points", f'=COUNTIF(Customers!C2:C{cl},">0")'), ("L7", "expired points", f"=SUM(Orders!O2:O{last})"),
    ("L8", "highest balance", f"=MAX(Customers!D2:D{cl})"),
    ("L8", "customer with highest balance", f"=INDEX(Customers!A2:A{cl},MATCH(MAX(Customers!D2:D{cl}),Customers!D2:D{cl},0))"),
    ("L8", "customers >= 4,000", f'=COUNTIF(Customers!D2:D{cl},">=4000")'),
    ("B1", "first order_id reaching 5,000", f"=MIN(Orders!P2:P{last})"),
    ("B1", "customer", f"=INDEX(Orders!C2:C{last},MATCH(MIN(Orders!P2:P{last}),Orders!A2:A{last},0))"),
    ("B2", "top month points", "=MAX(Months!B2:B13)"),
    ("B2", "top month", "=INDEX(Months!A2:A13,MATCH(MAX(Months!B2:B13),Months!B2:B13,0))"),
    ("B3", "customers ordering every month", f"=COUNTIF(Customers!F2:F{cl},12)"),
    ("B4", "Metro Mart rank", f'=INDEX(Customers!E2:E{cl},MATCH("Metro Mart",Customers!A2:A{cl},0))'),
]
Cs.cell(row=1, column=7, value="base_points").font = BOLD; Cs.cell(row=1, column=7).fill = HFILL
for i in range(2, cl + 1):
    Cs.cell(row=i, column=7, value=f"=SUMIFS(Orders!$G$2:$G${last},Orders!$C$2:$C${last},A{i})")
for a in ans: An.append(list(a))
An["C20"].number_format = "yyyy-mm"
sw.save(CH / "riverstone_rewards_solution.xlsx")
print("rewards case written")
