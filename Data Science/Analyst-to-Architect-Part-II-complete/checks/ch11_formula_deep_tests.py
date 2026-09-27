"""Chapter 11 (expanded): the formula toolkit sections. Every classic formula printed in sections 11.2, 11.3 and the
lookup extras is evaluated here by LibreOffice on ch11_practice.xlsx (and the messy workbook for INDIRECT).
Run from the book root:  python3 checks/ch11_formula_deep_tests.py"""
import subprocess, shutil, pathlib
from openpyxl import load_workbook
w = pathlib.Path("/tmp/ch11deep"); shutil.rmtree(w, ignore_errors=True); w.mkdir()
wb = load_workbook("companion/ch11/ch11_practice.xlsx")
T = wb.create_sheet("Tests")
R = lambda c: f"Sales!{c}2:{c}331"
V = f'{R("H")},"<>Cancelled"'
# a monthly summary block for rank / running total / OFFSET / INDEX column tests (rows 2..13, columns F:J)
T["F1"] = "month"; T["G1"] = "revenue"; T["H1"] = "rank"; T["I1"] = "running"; T["J1"] = "Q"
for i in range(12):
    r = i + 2
    T[f"F{r}"] = f"=DATE(2025,{i+1},1)"
    T[f"G{r}"] = f'=SUMIFS({R("J")},{R("K")},F{r},{V})'
    T[f"H{r}"] = f"=RANK.EQ(G{r},$G$2:$G$13,0)".replace("RANK.EQ", "_xlfn.RANK.EQ")
    T[f"I{r}"] = f'=SUMIFS({R("J")},{R("B")},"<="&EOMONTH(F{r},0),{V})'
    T[f"J{r}"] = f'=CHOOSE(ROUNDUP(MONTH(F{r})/3,0),"Q1","Q2","Q3","Q4")'
# a segment x quarter grid (rows 16..18, columns F:J) for INDEX(…,0,col)
T["F15"] = "segment"
for j, q in enumerate(["Q1", "Q2", "Q3", "Q4"]): T.cell(row=15, column=7 + j, value=q)
for i, s in enumerate(["Hospitality", "Retail", "Wholesale"]):
    T.cell(row=16 + i, column=6, value=s)
    for j in range(4):
        col = chr(ord("G") + j)
        T.cell(row=16 + i, column=7 + j, value=f'=SUMIFS({R("J")},{R("N")},$F{16+i},{R("L")},{col}$15,{V})')
T["K1"] = "Neha Kulkarni"
tests = [
 ("COUNTIF Retail all", f'=COUNTIF({R("N")},"Retail")'),
 ("COUNTIF valid", f'=COUNTIF({R("H")},"<>Cancelled")'),
 ("IF J5", '=IF(Sales!J5>=30000,"Large","Not large")'),
 ("IFS J2", '=_xlfn.IFS(Sales!J2>=30000,"Large",Sales!J2>=10000,"Medium",TRUE,"Small")'),
 ("ISTEXT quantities", f'=SUMPRODUCT(--ISTEXT({R("E")}))'),
 ("SUMIF Retail (all lines)", f'=SUMIF({R("N")},"Retail",{R("J")})'),
 ("SUMIFS Retail valid", f'=SUMIFS({R("J")},{R("N")},"Retail",{V})'),
 ("COUNTIF *Box*", f'=COUNTIF({R("P")},"*Box*")'),
 ("COUNTIF Storage Box ??L", f'=COUNTIF({R("P")},"Storage Box ??L")'),
 ("COUNTIFS rep non-blank", f'=COUNTIFS({R("I")},"<>")'),
 ("COUNTIFS rep blank", f'=COUNTIFS({R("I")},"")'),
 ("Q3 between dates", f'=SUMIFS({R("J")},{R("B")},">="&DATE(2025,7,1),{R("B")},"<="&DATE(2025,9,30),{V})'),
 ("OR via array constant", f'=SUM(SUMIFS({R("J")},{R("N")},{{"Retail","Hospitality"}},{V}))'),
 ("band 20k-50k count", f'=COUNTIFS({R("J")},">=20000",{R("J")},"<50000",{V})'),
 ("AVERAGEIFS wholesale line", f'=AVERAGEIFS({R("J")},{R("N")},"Wholesale",{V})'),
 ("distinct valid orders", f'=SUMPRODUCT(({R("H")}<>"Cancelled")/COUNTIFS({R("A")},{R("A")}))'),
 ("AOV", f'=SUMIFS({R("J")},{V})/SUMPRODUCT(({R("H")}<>"Cancelled")/COUNTIFS({R("A")},{R("A")}))'),
 ("AOV wholesale", f'=SUMIFS({R("J")},{R("N")},"Wholesale",{V})/SUMPRODUCT(({R("N")}="Wholesale")*({R("H")}<>"Cancelled")/COUNTIFS({R("A")},{R("A")}))'),
 ("MAXIFS retail line", f'=_xlfn.MAXIFS({R("J")},{R("N")},"Retail",{V})'),
 ("MAXIFS latest wholesale date", f'=TEXT(_xlfn.MAXIFS({R("B")},{R("N")},"Wholesale"),"yyyy-mm-dd")'),
 ("MINIFS smallest discount >0", f'=_xlfn.MINIFS({R("G")},{R("G")},">0")'),
 ("weighted avg discount valid", f'=SUMPRODUCT(({R("H")}<>"Cancelled")*{R("E")}*{R("F")}*{R("G")})/SUMPRODUCT(({R("H")}<>"Cancelled")*{R("E")}*{R("F")})'),
 ("simple avg discount valid", f'=AVERAGEIFS({R("G")},{V})'),
 ("SUMPRODUCT MONTH=11", f'=SUMPRODUCT((MONTH({R("B")})=11)*({R("H")}<>"Cancelled")*{R("J")})'),
 ("OR across columns revenue", f'=SUMPRODUCT(((({R("N")}="Wholesale")+({R("E")}>=50))>0)*({R("H")}<>"Cancelled")*{R("J")})'),
 ("OR double count (wrong)", f'=COUNTIFS({R("N")},"Wholesale",{V})+COUNTIFS({R("E")},">=50",{V})'),
 ("OR count (right)", f'=SUMPRODUCT(((({R("N")}="Wholesale")+({R("E")}>=50))>0)*({R("H")}<>"Cancelled"))'),
 ("distinct customers", f'=SUMPRODUCT(1/COUNTIF({R("M")},{R("M")}))'),
 ("rank Oct", "=H11"), ("rank Jun", "=H7"), ("running Sep", "=I10"), ("running Dec", "=I13"), ("CHOOSE quarter May", "=J6"),
 ("OFFSET last 3 months", "=SUM(OFFSET(G2,COUNT(G2:G13)-3,0,3,1))"),
 ("INDEX column Q2 total", '=SUM(INDEX(G16:J18,0,MATCH("Q2",G15:J15,0)))'),
 ("INDEX row Wholesale total", '=SUM(INDEX(G16:J18,MATCH("Wholesale",F16:F18,0),0))'),
 ("SWITCH row 2", '=_xlfn.SWITCH(Sales!H2,"Delivered","Closed","Shipped","In transit","Pending","Open","Cancelled","Void")'),
 ("IFNA MATCH", '=_xlfn.IFNA(MATCH("9999",Customers!A2:A25,0),"not found")'),
 ("REPT bar", '=REPT("|",ROUND(0.393*20,0))'),
 ("SEARCH box", '=SEARCH("box","Lunch Box Set")'), ("FIND box", '=IFERROR(FIND("box","Lunch Box Set"),"#VALUE!")'),
 ("EXACT", '=EXACT("0014","0014")'),
 ("EDATE Jan31+1", '=TEXT(EDATE(DATE(2025,1,31),1),"yyyy-mm-dd")'),
 ("WORKDAY.INTL Sun-only", '=TEXT(WORKDAY.INTL(DATE(2025,12,24),3,11),"yyyy-mm-dd")'),
 ("WORKDAY.INTL + holiday", '=TEXT(WORKDAY.INTL(DATE(2025,12,24),3,11,DATE(2025,12,25)),"yyyy-mm-dd")'),
 ("NETWORKDAYS.INTL Dec Sun-only", "=NETWORKDAYS.INTL(DATE(2025,12,1),DATE(2025,12,31),11)"),
 ("WEEKNUM Jan 2", "=WEEKNUM(DATE(2025,1,2))"), ("ISOWEEKNUM Dec 29", "=_xlfn.ISOWEEKNUM(DATE(2025,12,29))"), ("WEEKNUM Dec 29 (LibreOffice differs from Excel; Excel type 1 = 53, computed by hand)", "=WEEKNUM(DATE(2025,12,29),1)"),
 ("DATEDIF m Sharma", '=DATEDIF(DATE(2025,1,24),DATE(2025,12,5),"m")'), ("DATEDIF d Sharma", '=DATEDIF(DATE(2025,1,24),DATE(2025,12,5),"d")'),
 ("LARGE 2", f'=LARGE({R("J")},2)'), ("SMALL 1", f'=SMALL({R("J")},1)'),
 ("RANK.EQ 51520", f'=_xlfn.RANK.EQ(51520,{R("J")},0)'), ("COUNTIF >51520 +1", f'=COUNTIF({R("J")},">"&51520)+1'),
 ("PERCENTILE 0.9", f'=_xlfn.PERCENTILE.INC({R("J")},0.9)'), ("QUARTILE 1", f'=_xlfn.QUARTILE.INC({R("J")},1)'), ("MEDIAN", f'=MEDIAN({R("J")})'),
 ("MODE quantity", f'=_xlfn.MODE.SNGL({R("E")})'),
 ("MROUND 47,10", "=MROUND(47,10)"), ("CEILING.MATH 47,10", "=_xlfn.CEILING.MATH(47,10)"), ("FLOOR.MATH 47,10", "=_xlfn.FLOOR.MATH(47,10)"),
 ("cartons for 45", "=_xlfn.CEILING.MATH(45,10)/10"),
 ("PMT van loan", "=PMT(10%/12,36,-500000)"),
 ("AGGREGATE ignore errors", '=_xlfn.AGGREGATE(9,6,Tests!G2:G13)'),
 ("TEXTJOIN wholesale (array)", '=_xlfn.TEXTJOIN(", ",TRUE,IF(Customers!D2:D25="Wholesale",Customers!B2:B25,""))'),
 ("LEFT/FIND first name", '=LEFT(K1,FIND(" ",K1)-1)'),
 ("INDEX/MATCH multi-criteria", f'=INDEX({R("A")},MATCH(1,({R("M")}="Harbour Traders")*({R("P")}="Industrial Crate"),0))'),
 ("count Harbour crates lines", f'=COUNTIFS({R("M")},"Harbour Traders",{R("P")},"Industrial Crate")'),
 ("PERCENTOF-style share wholesale", f'=SUMIFS({R("J")},{R("N")},"Wholesale",{V})/SUMIFS({R("J")},{V})'),
]
for i, (n, f) in enumerate(tests, start=1):
    T.cell(row=i + 20, column=1, value=n); c = T.cell(row=i + 20, column=2, value=f)
wb.save(w / "t.xlsx")
subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(w / "o"), str(w / "t.xlsx")], capture_output=True)
r = load_workbook(w / "o" / "t.xlsx", data_only=True)["Tests"]
for i, (n, f) in enumerate(tests, start=1):
    print(f"{n:34} -> {r.cell(row=i + 20, column=2).value!r}")
print("grid:", [[r.cell(row=16 + i, column=7 + j).value for j in range(4)] for i in range(3)])
# INDIRECT across month sheets of the messy workbook
mw = load_workbook("companion/ch11/month_end_pack_2025_messy.xlsx")
S2 = mw["Summary"]
for m in range(12):
    S2.cell(row=4 + m, column=7, value=f"=SUMIFS(INDIRECT(\"'\"&A{4+m}&\"'!J2:J100\"),INDIRECT(\"'\"&A{4+m}&\"'!H2:H100\"),\"<>Cancelled\")")
mw.save(w / "m.xlsx")
subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(w / "o"), str(w / "m.xlsx")], capture_output=True)
rm = load_workbook(w / "o" / "m.xlsx", data_only=True)["Summary"]
print("INDIRECT month totals:", [rm.cell(row=4 + m, column=7).value for m in range(12)])
