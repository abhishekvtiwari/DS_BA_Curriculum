"""Chapter 10: every classic formula printed in the chapter, evaluated by LibreOffice on the Chapter 10 data.
Run from the book root after companion/ch10/build_ch10_files.py:  python3 checks/ch10_formula_tests.py
XLOOKUP is evaluated with the equivalent INDEX/MATCH (LibreOffice 24.2 can't evaluate XLOOKUP); see manual checks."""
import subprocess, shutil, pathlib, json
from openpyxl import load_workbook
src = pathlib.Path("companion/ch10/ch10_tracker_check.xlsx")
work = pathlib.Path("/tmp/ch10ft"); shutil.rmtree(work, ignore_errors=True); work.mkdir()
wb = load_workbook(src)
det = load_workbook("companion/ch10/ch10_practice.xlsx")["Cell detective"]
cd = wb.create_sheet("Cell detective")
for row in det.iter_rows(values_only=True): cd.append(row)
cd["A5"].number_format = "yyyy-mm-dd"
T = wb.create_sheet("Tests")
tests = [
 ("COUNT order_id", "=COUNT(Data!A2:A331)"),
 ("COUNTA sales_rep", "=COUNTA(Data!I2:I331)"),
 ("COUNTBLANK sales_rep", "=COUNTBLANK(Data!I2:I331)"),
 ("SUM net all", "=SUM(Data!J2:J331)"),
 ("SUMIFS non-cancelled", '=SUMIFS(Data!J2:J331,Data!H2:H331,"<>Cancelled")'),
 ("AVERAGE all", "=AVERAGE(Data!J2:J331)"),
 ("AVERAGEIFS non-cancelled", '=AVERAGEIFS(Data!J2:J331,Data!H2:H331,"<>Cancelled")'),
 ("MAX", "=MAX(Data!J2:J331)"), ("MIN", "=MIN(Data!J2:J331)"),
 ("COUNTIF disc>0", '=COUNTIF(Data!G2:G331,">0")'),
 ("COUNTIFS Delivered disc>0", '=COUNTIFS(Data!H2:H331,"Delivered",Data!G2:G331,">0")'),
 ("COUNTIF >=50000", '=COUNTIF(Data!J2:J331,">=50000")'),
 ("IF row5", '=IF(Data!J5>=50000,"Big","Normal")'),
 ("ROUND J5", "=ROUND(Data!J5,0)"),
 ("TEXT B2 mmm yyyy", '=TEXT(Data!B2,"mmm yyyy")'),
 ("EOMONTH B4", '=TEXT(EOMONTH(Data!B4,0),"yyyy-mm-dd")'),
 ("days to month end", "=EOMONTH(Data!B2,0)-Data!B2"),
 ("NETWORKDAYS Dec", "=NETWORKDAYS(DATE(2025,12,1),DATE(2025,12,31))"),
 ("serial B2", "=Data!B2*1"), ("WEEKDAY B2", "=WEEKDAY(Data!B2,2)"),
 ("TEXT 5 0000", '=TEXT(5,"0000")'),
 ("ISNUMBER A2", "=ISNUMBER('Cell detective'!A2)"), ("ISTEXT A3", "=ISTEXT('Cell detective'!A3)"),
 ("LEN A3", "=LEN('Cell detective'!A3)"), ("LEN A4", "=LEN('Cell detective'!A4)"),
 ("ISNUMBER A5 date", "=ISNUMBER('Cell detective'!A5)"), ("ISNUMBER A6 textdate", "=ISNUMBER('Cell detective'!A6)"),
 ("SUM A2:A4", "=SUM('Cell detective'!A2:A4)"),
 ("VALUE TRIM A4", "=VALUE(TRIM('Cell detective'!A4))"),
 ("A7 displayed", '=TEXT(\'Cell detective\'!A7,"#,##0")'),
 ("PROPER", '=PROPER("METRO MART")'), ("LEFT", '=LEFT("0005",2)'), ("RIGHT", '=RIGHT("0005",1)'),
 ("MID", '=MID("RS-2025-0418",4,4)'),
 ("SUBSTITUTE", '=SUBSTITUTE("Storage Box 10L","Box","Bin")'),
 ("join", '=Customers!B3&" ("&Customers!A3&")"'),
 ("share correct row3", "=Data!J3/SUM(Data!$J$2:$J$331)"),
 ("share relative-copy row3", "=Data!J3/SUM(Data!J3:J332)"),
 ("share relative-copy row331", "=Data!J331/SUM(Data!J331:J660)"),
 ("mixed grid 750x5%", "=750*(1-0.05)"), ("mixed grid 1400x10%", "=1400*(1-0.1)"),
 ("Nov revenue", '=SUMIFS(Data!J2:J331,Data!B2:B331,">="&DATE(2025,11,1),Data!B2:B331,"<="&DATE(2025,11,30),Data!H2:H331,"<>Cancelled")'),
 ("Nov wholesale", '=SUMIFS(Data!J2:J331,Data!B2:B331,">="&DATE(2025,11,1),Data!B2:B331,"<="&DATE(2025,11,30),Data!H2:H331,"<>Cancelled",Data!N2:N331,"Wholesale")'),
 ("lookup 0005", '=INDEX(Customers!B2:B25,MATCH("0005",Customers!A2:A25,0))'),
 ("lookup number 5", '=IFERROR(INDEX(Customers!B2:B25,MATCH(5,Customers!A2:A25,0)),"not found")'),
 ("IFERROR", '=IFERROR(1/0,"n/a")'),
 ("first-line flag sum", '=SUMIFS(Data!L2:L331,Data!H2:H331,"<>Cancelled")'),
 ("IFS", '=IFS(Data!J5>=50000,"Large",Data!J5>=10000,"Medium",TRUE,"Small")'),
 ("AND", '=AND(Data!H5="Delivered",Data!G5>0)'), ("OR", '=OR(Data!G2>=10,Data!E2>=50)'),
 ("row count Farah", '=COUNTIF(Data!I2:I331,"Farah Khan")'),
 ("wildcard Hotels", '=COUNTIF(Customers!B2:B25,"*Hotel*")'),
]
for i,(name,f) in enumerate(tests, start=1):
    T.cell(row=i, column=1, value=name)
    T.cell(row=i, column=2, value=f.replace("=IFS(", "=_xlfn.IFS("))
wb.save(work/"t.xlsx")
subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",str(work/"out"),str(work/"t.xlsx")],capture_output=True)
r = load_workbook(work/"out"/"t.xlsx", data_only=True)["Tests"]
for i,(name,f) in enumerate(tests, start=1):
    print(f"{name:28} {f:60.60} -> {r.cell(row=i,column=2).value!r}")
